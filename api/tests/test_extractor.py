import io

import docx
import fitz
import openpyxl
import pytest

from services.extractor import UnsupportedFormatError, extract_text


def _make_pdf_bytes(text: str) -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text)
    return doc.tobytes()


def _make_docx_bytes(text: str) -> bytes:
    document = docx.Document()
    document.add_paragraph(text)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _make_xlsx_bytes(text: str) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = text
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_extract_pdf():
    content = _make_pdf_bytes("Pregunta de prueba PDF")
    result = extract_text(content, "test.pdf")
    assert "Pregunta de prueba PDF" in result


def test_extract_docx():
    content = _make_docx_bytes("Pregunta de prueba DOCX")
    result = extract_text(content, "test.docx")
    assert "Pregunta de prueba DOCX" in result


def test_extract_xlsx():
    content = _make_xlsx_bytes("Pregunta de prueba XLSX")
    result = extract_text(content, "test.xlsx")
    assert "Pregunta de prueba XLSX" in result


def test_unsupported_format():
    with pytest.raises(UnsupportedFormatError):
        extract_text(b"data", "test.txt")
